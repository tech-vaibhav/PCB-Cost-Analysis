"""
test_3d_parser.py  —  smoke test + Excel/CSV export for the 3D parser service.

Usage:
    .\\venv\\Scripts\\python.exe test_3d_parser.py path/to/your_model.stp
"""
import sys
import csv
import json
import os
from pathlib import Path

# Make sure the project root is on sys.path
sys.path.insert(0, str(Path(__file__).parent))

from backend.services.printer_3d_parser import parse_3d_file

RESULTS_DIR = Path(__file__).parent / "backend" / "results"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _flatten(prefix: str, obj, rows: list):
    """Recursively flatten nested dicts/lists into (key, value) rows."""
    if isinstance(obj, dict):
        for k, v in obj.items():
            _flatten(f"{prefix}{k}." if prefix else f"{k}.", v, rows)
    elif isinstance(obj, list):
        for i, item in enumerate(obj):
            _flatten(f"{prefix}{i}.", item, rows)
    else:
        rows.append((prefix.rstrip("."), obj))


def _write_csv(stem: str, rows: list[tuple]) -> Path:
    RESULTS_DIR.mkdir(exist_ok=True)
    path = RESULTS_DIR / f"3d_{stem}.csv"
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Parameter", "Value"])
        w.writerows(rows)
    return path


def _write_excel(stem: str, rows: list[tuple], data: dict) -> Path:
    try:
        import openpyxl
        from openpyxl.styles import (
            Font, PatternFill, Alignment, Border, Side
        )
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  openpyxl not installed — skipping Excel export.")
        print("  Run: .\\venv\\Scripts\\pip.exe install openpyxl")
        return None

    RESULTS_DIR.mkdir(exist_ok=True)
    path = RESULTS_DIR / f"3d_{stem}.xlsx"
    wb = openpyxl.Workbook()

    # ── Sheet 1: Flat parameter table ────────────────────────────────────────
    ws = wb.active
    ws.title = "All Parameters"

    # Styles
    hdr_fill   = PatternFill("solid", fgColor="1F3864")
    hdr_font   = Font(bold=True, color="FFFFFF", size=11)
    alt_fill   = PatternFill("solid", fgColor="DCE6F1")
    none_fill  = PatternFill("solid", fgColor="FFF2CC")
    sec_fill   = PatternFill("solid", fgColor="2E75B6")
    sec_font   = Font(bold=True, color="FFFFFF", size=10)
    thin       = Side(style="thin", color="AAAAAA")
    border     = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Header row
    ws.append(["Parameter", "Value", "Unit / Note"])
    for col in range(1, 4):
        c = ws.cell(1, col)
        c.fill = hdr_fill
        c.font = hdr_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    ws.row_dimensions[1].height = 22

    # Section grouping
    SECTIONS = {
        "file_info":    ("FILE INFO",    ""),
        "dimensions":   ("DIMENSIONS",   "mm"),
        "geometry":     ("GEOMETRY",     ""),
        "mesh_stats":   ("TOPOLOGY",     "count"),
        "raw_metadata": ("RAW METADATA", ""),
        "warnings":     ("WARNINGS",     ""),
    }

    row_idx = 2
    for section_key, (section_label, unit_hint) in SECTIONS.items():
        section_data = data.get(section_key)
        if section_data is None:
            continue

        # Section header
        ws.merge_cells(f"A{row_idx}:C{row_idx}")
        c = ws.cell(row_idx, 1, section_label)
        c.fill = sec_fill
        c.font = sec_font
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[row_idx].height = 18
        row_idx += 1

        # Flatten section rows
        sec_rows: list[tuple] = []
        _flatten("", section_data, sec_rows)

        for i, (k, v) in enumerate(sec_rows):
            fill = alt_fill if i % 2 == 0 else PatternFill()
            val_fill = none_fill if v is None else fill

            unit_cell = ""
            k_lower = k.lower()
            if "mm" in k_lower:
                unit_cell = "mm"
            elif "cm3" in k_lower or "volume" in k_lower:
                unit_cell = "cm³"
            elif "cm2" in k_lower or "area" in k_lower:
                unit_cell = "cm²"
            elif "bytes" in k_lower:
                unit_cell = "bytes"
            elif "count" in k_lower or k_lower in ("solid_count","shell_count","face_count","edge_count","vertex_count","compound_parts"):
                unit_cell = "count"

            for col, val in enumerate([k, v if v is not None else "—", unit_cell], start=1):
                c = ws.cell(row_idx, col, val)
                c.fill = val_fill if col == 2 else fill
                c.border = border
                c.alignment = Alignment(vertical="center", indent=1)
                if col == 1:
                    c.font = Font(bold=True, size=10)
                else:
                    c.font = Font(size=10)
            ws.row_dimensions[row_idx].height = 17
            row_idx += 1

    # Column widths
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 14
    ws.freeze_panes = "A2"

    # ── Sheet 2: Summary card ─────────────────────────────────────────────────
    ws2 = wb.create_sheet("Summary")
    ws2.sheet_view.showGridLines = False

    title_font = Font(bold=True, size=16, color="1F3864")
    label_font = Font(bold=True, size=11, color="2E75B6")
    val_font   = Font(size=11)

    ws2["A1"] = f"3D File Analysis — {data['file_info']['filename']}"
    ws2["A1"].font = title_font
    ws2.row_dimensions[1].height = 30

    summary_rows = [
        ("", ""),
        ("FILE", ""),
        ("Filename",    data["file_info"]["filename"]),
        ("Format",      data["file_info"]["format"]),
        ("File Size",   f"{data['file_info']['size_bytes'] / 1024:.1f} KB"),
        ("", ""),
        ("DIMENSIONS", ""),
        ("Length (X)",  f"{data['dimensions']['length_mm']} mm"),
        ("Width  (Y)",  f"{data['dimensions']['width_mm']} mm"),
        ("Height (Z)",  f"{data['dimensions']['height_mm']} mm"),
        ("BBox Xmin",   f"{data['dimensions']['bbox_xmin']} mm"),
        ("BBox Xmax",   f"{data['dimensions']['bbox_xmax']} mm"),
        ("BBox Ymin",   f"{data['dimensions']['bbox_ymin']} mm"),
        ("BBox Ymax",   f"{data['dimensions']['bbox_ymax']} mm"),
        ("BBox Zmin",   f"{data['dimensions']['bbox_zmin']} mm"),
        ("BBox Zmax",   f"{data['dimensions']['bbox_zmax']} mm"),
        ("", ""),
        ("GEOMETRY", ""),
        ("Volume",      f"{data['geometry']['volume_cm3']} cm³"),
        ("Surface Area",f"{data['geometry']['surface_area_cm2']} cm²"),
        ("", ""),
        ("TOPOLOGY", ""),
        ("Solid Count",    data["mesh_stats"]["solid_count"]),
        ("Shell Count",    data["mesh_stats"]["shell_count"]),
        ("Face Count",     data["mesh_stats"]["face_count"]),
        ("Edge Count",     data["mesh_stats"]["edge_count"]),
        ("Vertex Count",   data["mesh_stats"]["vertex_count"]),
        ("Compound Parts", data["mesh_stats"]["compound_parts"]),
        ("Shape Type",     data["raw_metadata"].get("shape_type", "—")),
    ]

    if data["warnings"]:
        summary_rows += [("", ""), ("WARNINGS", "")]
        for w in data["warnings"]:
            summary_rows.append(("  ⚠", w))

    for r_offset, (label, value) in enumerate(summary_rows, start=2):
        if label in ("FILE", "DIMENSIONS", "GEOMETRY", "TOPOLOGY", "WARNINGS"):
            c = ws2.cell(r_offset, 1, label)
            c.font = Font(bold=True, size=12, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2E75B6")
            ws2.merge_cells(f"A{r_offset}:B{r_offset}")
            ws2.row_dimensions[r_offset].height = 20
        elif label == "":
            ws2.row_dimensions[r_offset].height = 8
        else:
            ws2.cell(r_offset, 1, label).font = label_font
            ws2.cell(r_offset, 2, value).font = val_font
            ws2.row_dimensions[r_offset].height = 17

    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 28

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    if len(sys.argv) < 2:
        print("Usage: python test_3d_parser.py <path_to_stp_file>")
        sys.exit(1)

    file_path = Path(sys.argv[1])
    if not file_path.exists():
        print(f"File not found: {file_path}")
        sys.exit(1)

    print(f"\nParsing: {file_path.name}  ({file_path.stat().st_size / 1024:.1f} KB)")
    print("-" * 60)

    raw = file_path.read_bytes()
    result = parse_3d_file(raw, filename=file_path.name)
    data   = result.model_dump()

    # ── Flatten all fields ────────────────────────────────────────────────────
    all_rows: list[tuple] = []
    _flatten("", data, all_rows)

    # ── CSV ───────────────────────────────────────────────────────────────────
    stem     = file_path.stem
    csv_path = _write_csv(stem, all_rows)
    print(f"\nCSV  saved -> {csv_path}")

    # ── Excel ─────────────────────────────────────────────────────────────────
    xl_path = _write_excel(stem, all_rows, data)
    if xl_path:
        print(f"XLSX saved -> {xl_path}")
        # Auto-open the Excel file
        os.startfile(str(xl_path))

    # ── Console summary ───────────────────────────────────────────────────────
    print("\n" + "=" * 60)
    print("FILE")
    fi = data["file_info"]
    print(f"  Format : {fi['format']}")
    print(f"  Size   : {fi['size_bytes'] / 1024:.1f} KB")

    print("\nDIMENSIONS (bounding box)")
    d = data["dimensions"]
    print(f"  L x W x H : {d['length_mm']} x {d['width_mm']} x {d['height_mm']} mm")

    print("\nGEOMETRY")
    g = data["geometry"]
    print(f"  Volume       : {g['volume_cm3']} cm3")
    print(f"  Surface Area : {g['surface_area_cm2']} cm2")

    print("\nTOPOLOGY")
    m = data["mesh_stats"]
    print(f"  Solids   : {m['solid_count']}")
    print(f"  Shells   : {m['shell_count']}")
    print(f"  Faces    : {m['face_count']}")
    print(f"  Edges    : {m['edge_count']}")
    print(f"  Vertices : {m['vertex_count']}")
    print(f"  Parts    : {m['compound_parts']}")

    if data["warnings"]:
        print("\nWARNINGS")
        for w in data["warnings"]:
            print(f"  * {w}")

    print("=" * 60)


if __name__ == "__main__":
    main()
